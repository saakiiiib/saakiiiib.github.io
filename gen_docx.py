"""Bundle the portfolio source code into a single Word (.docx) document."""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_LINE_SPACING

HERE = os.path.dirname(os.path.abspath(__file__))
FILES = ["index.html", "resume.html", "gen_figs.py", "render_field.py"]

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

title = doc.add_heading("Engineering Research Portfolio — Source Code", level=0)
p = doc.add_paragraph(
    "Generated bundle of all source files. "
    "Project folder: web portofolio (Desktop). "
    "PERSONALIZE: search for \u201cAlex Rivera\u201d, \u201chello@alexrivera.dev\u201d and \u201cTODO\u201d "
    "in index.html and resume.html."
)
p.runs[0].font.size = Pt(10)
doc.add_paragraph("Files: " + ", ".join(FILES) + ".")

for name in FILES:
    path = os.path.join(HERE, name)
    if not os.path.exists(path):
        doc.add_heading("MISSING: " + name, level=1)
        continue
    with open(path, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc.add_page_break()
    doc.add_heading(name, level=1)
    meta = doc.add_paragraph(name + "  —  " + str(os.path.getsize(path)) + " bytes")
    meta.runs[0].font.size = Pt(9)
    meta.runs[0].font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    doc.add_paragraph()

    for line in lines:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = para.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

out = os.path.join(HERE, "portfolio_code.docx")
try:
    doc.save(out)
except PermissionError:
    out = os.path.join(HERE, "portfolio_code_v2.docx")
    doc.save(out)
print("saved", out, "->", os.path.getsize(out), "bytes")
